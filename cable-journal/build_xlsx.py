"""Cable journal workbook: journal, summary by cable, explanations. Usage: python3 build_xlsx.py out.xlsx"""
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins

from config import DROPS, GROUPS_JSON, LENGTH_FACTOR, NOTES_JSON, r2

OUT = sys.argv[1]
FACTOR = f'*{LENGTH_FACTOR:g}' if LENGTH_FACTOR != 1 else ''   # every length is multiplied by it
groups = json.load(open(GROUPS_JSON))
notes = json.load(open(NOTES_JSON))

F = 'Arial'
f_title = Font(name=F, size=14, bold=True)
f_sub = Font(name=F, size=9, italic=True, color='404040')
f_hdr = Font(name=F, size=9, bold=True)
f_num = Font(name=F, size=8, color='606060')
f_txt = Font(name=F, size=10)
f_in = Font(name=F, size=10, color='0000FF')          # hardcoded inputs (from drawing / customer)
f_link = Font(name=F, size=10, color='008000')        # links to another sheet
f_bold = Font(name=F, size=10, bold=True)
fill_hdr = PatternFill('solid', fgColor='DCE6F1')
fill_sec = PatternFill('solid', fgColor='EFEFEF')
fill_sub = PatternFill('solid', fgColor='F7F7F7')
fill_tot = PatternFill('solid', fgColor='E2EFDA')
fill_ass = PatternFill('solid', fgColor='FFF2CC')        # assumptions to verify
fill_par = PatternFill('solid', fgColor='FFFF99')        # user-editable parameters
thin = Side(style='thin', color='808080')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_c = Alignment(horizontal='center', vertical='center', wrap_text=True)
wrap_l = Alignment(horizontal='left', vertical='center', wrap_text=True)
right = Alignment(horizontal='right', vertical='center')
center = Alignment(horizontal='center', vertical='center')


# ---- column map of the journal ----
C_NO, C_DES, C_FROM, C_TO, C_CONS, C_BRAND, C_CORES = 'A', 'B', 'C', 'D', 'E', 'F', 'G'
C_H, C_NP, C_NS, C_NW, C_NL, C_DROPS, C_LEN, C_NOTE = 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O'
LAST_COL = 15
COUNT_COLS = C_NP + C_NS + C_NW + C_NL           # 'IJKL'
LEN_COLS = C_H + C_DROPS + C_LEN                  # 'HMN'
SUM_COLS = C_H + COUNT_COLS + C_DROPS + C_LEN     # 'HIJKLMN'

wb = Workbook()
ws = wb.active
ws.title = 'Кабельный журнал'
JS = "'Кабельный журнал'"

widths = {'A': 5, 'B': 11, 'C': 24, 'D': 22, 'E': 36, 'F': 15, 'G': 10, 'H': 15, 'I': 8, 'J': 10, 'K': 13,
          'L': 13, 'M': 10, 'N': 11, 'O': 58}
for c, w in widths.items():
    ws.column_dimensions[c].width = w

ws['A1'] = 'КАБЕЛЬНЫЙ ЖУРНАЛ'
ws['A1'].font = f_title
ws.merge_cells('A1:O1')
ws['A2'] = ('Составлен по присланному чертежу DWG (план этажа). Длина кабеля = горизонтальная трасса + спуски, в метрах. '
            'Номера групп условные (на чертеже их нет) и совпадают со «Схемой групп» (PDF); названия помещений — по данным заказчика.')
ws['A2'].font = f_sub
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:O2')
ws.row_dimensions[2].height = 26

# ---- parameters (customer inputs) ----
ws['C3'] = 'Исходные данные (можно менять)'
ws['C3'].font = f_bold
params = [
    ('Спуск к щитку, м', DROPS['panel'], 'задано заказчиком; 1 спуск на каждую линию, выходящую из щитка'),
    ('Спуск к розетке, м', DROPS['socket'], 'задано заказчиком; 1 спуск на каждую розетку'),
    ('Спуск к выключателю, м', DROPS['switch'], 'задано заказчиком; 1 спуск на каждый выключатель'),
    ('Спуск к светильнику, м', DROPS['luminaire'],
     'задано заказчиком; 1 спуск от каждой коробки-светильника (квадрат с диагональю)'),
]
P0 = 4
for i, (lab, val, note) in enumerate(params):
    r = P0 + i
    ws.cell(r, 3, lab).font = f_txt
    c = ws.cell(r, 4, val)
    c.font = f_in; c.fill = fill_par; c.border = border; c.alignment = Alignment(horizontal='center')
    c.number_format = '0.00'
    ws.cell(r, 5, note).font = f_sub
P_PANEL, P_SOCK, P_SW, P_LUM = (f'$D${P0 + i}' for i in range(4))
DROP_COLS = ((C_NP, P_PANEL), (C_NS, P_SOCK), (C_NW, P_SW), (C_NL, P_LUM))

ws['H4'] = 'Условные обозначения:'
ws['H4'].font = f_bold
leg = [('Синий шрифт', 'данные с чертежа и исходные данные', f_in, None),
       ('Жёлтая заливка', 'принятые допущения, проверьте (марка и сечение на чертеже не указаны)', f_txt, fill_ass),
       ('Чёрный шрифт', 'формулы', f_txt, None)]
for i, (a, b, fo, fi) in enumerate(leg):
    c = ws.cell(5 + i, 8, a); c.font = fo
    if fi:
        c.fill = fi; ws.cell(5 + i, 9).fill = fi
    ws.merge_cells(start_row=5 + i, start_column=8, end_row=5 + i, end_column=9)
    ws.cell(5 + i, 10, b).font = f_sub

# ---- table header ----
H1, H2, HN = 10, 11, 12
hdr = [
    (C_NO, '№ п/п', None), (C_DES, 'Обозначение кабеля', None), (C_FROM, 'Трасса', 'Начало'), (C_TO, None, 'Конец'),
    (C_CONS, 'Потребители', None), (C_BRAND, 'Кабель', 'Марка'), (C_CORES, None, 'Кол-во и сечение жил, мм²'),
    (C_H, 'Горизонтальная трасса, м', None), (C_NP, 'Количество спусков, шт.', 'к щитку'),
    (C_NS, None, 'к розеткам'), (C_NW, None, 'к выключателям'), (C_NL, None, 'к светильникам'),
    (C_DROPS, 'Спуски, м', None), (C_LEN, 'Длина кабеля, м', None), (C_NOTE, 'Примечание', None),
]
for col, top, bot in hdr:
    for r in (H1, H2):
        c = ws[f'{col}{r}']
        c.font = f_hdr; c.fill = fill_hdr; c.border = border; c.alignment = wrap_c
    if top and bot:
        ws[f'{col}{H1}'] = top; ws[f'{col}{H2}'] = bot
    elif top:
        ws[f'{col}{H1}'] = top
        ws.merge_cells(f'{col}{H1}:{col}{H2}')
    else:
        ws[f'{col}{H2}'] = bot
ws.merge_cells(f'{C_FROM}{H1}:{C_TO}{H1}')
ws.merge_cells(f'{C_BRAND}{H1}:{C_CORES}{H1}')
ws.merge_cells(f'{C_NP}{H1}:{C_NL}{H1}')
ws.row_dimensions[H1].height = 30
ws.row_dimensions[H2].height = 40
for i in range(LAST_COL):
    c = ws.cell(HN, i + 1, i + 1)
    c.font = f_num; c.alignment = wrap_c; c.border = border
ws.freeze_panes = f'A{HN + 1}'


def est_height(texts_widths, base=13.0):
    lines = 1
    for t, w in texts_widths:
        if not t:
            continue
        per = max(1, int(w * 1.15))
        n = sum(max(1, math.ceil(len(part) / per)) for part in str(t).split('\n'))
        lines = max(lines, n)
    return max(15, lines * base + 3)


def box_row(r, fill=None):
    for col in range(1, LAST_COL + 1):
        cell = ws.cell(r, col)
        cell.border = border
        if fill:
            cell.fill = fill


row = HN + 1
subtot_rows = {}
sec_bounds = {}
sections = ['Розеточная сеть', 'Кондиционеры', 'Рабочее освещение', 'Аварийное освещение']
n_pp = 0
for sec in sections:
    items = [g for g in groups if g['system'] == sec]
    ws.cell(row, 1, sec).font = f_bold
    box_row(row, fill_sec)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_COL)
    row += 1
    first = row
    for g in items:
        n_pp += 1
        vals = {
            C_NO: n_pp, C_DES: g['des'], C_FROM: g['start'], C_TO: g['end'], C_CONS: g['consumers'],
            C_BRAND: g['cable'], C_CORES: g['cores'], C_H: r2(g['horiz'] * LENGTH_FACTOR),
            C_NP: g['n_panel'], C_NS: g['n_sock'], C_NW: g['n_sw'], C_NL: g['n_lum'],
            C_DROPS: '=' + '+'.join(f'ROUND({cnt}{row}*{par}{FACTOR},2)' for cnt, par in DROP_COLS),
            C_LEN: f'={C_H}{row}+{C_DROPS}{row}',
            C_NOTE: g['note'] or None,
        }
        for col, v in vals.items():
            c = ws[f'{col}{row}']
            c.value = v
            c.border = border
            c.font = f_in if col in C_H + COUNT_COLS else f_txt
            if col in C_BRAND + C_CORES:
                c.fill = fill_ass
            if col in C_NO + C_DES + C_CORES + COUNT_COLS:
                c.alignment = center
            elif col in LEN_COLS:
                c.alignment = right
            else:
                c.alignment = wrap_l
            if col in LEN_COLS:
                c.number_format = '0.00'
        ws[f'{C_LEN}{row}'].font = Font(name=F, size=10, bold=True)
        ws.row_dimensions[row].height = est_height([(g['start'], widths[C_FROM]), (g['end'], widths[C_TO]),
                                                   (g['consumers'], widths[C_CONS]), (g['note'], widths[C_NOTE])])
        row += 1
    last = row - 1
    sec_bounds[sec] = (first, last)
    ws.cell(row, 5, f'Итого: {sec.lower()}').font = f_bold
    for col in SUM_COLS:
        c = ws[f'{col}{row}']
        c.value = f'=SUM({col}{first}:{col}{last})'
        c.font = f_bold
        c.number_format = '0.00' if col in LEN_COLS else '0'
        c.alignment = right if col in LEN_COLS else center
    box_row(row, fill_sub)
    subtot_rows[sec] = row
    row += 1

TOT = row
ws.cell(TOT, 5, 'ВСЕГО по журналу').font = Font(name=F, size=11, bold=True)
for col in SUM_COLS:
    c = ws[f'{col}{TOT}']
    c.value = '=' + '+'.join(f'{col}{r}' for r in subtot_rows.values())
    c.font = Font(name=F, size=11, bold=True)
    c.number_format = '0.00' if col in LEN_COLS else '0'
    c.alignment = right if col in LEN_COLS else center
box_row(TOT, fill_tot)
FIRST_DATA, LAST_DATA = HN + 1, TOT - 1
ws.cell(TOT + 2, 3, 'Методика расчёта, допущения и замечания к чертежу — на листе «Пояснения». '
                    'Итоги по маркам кабеля — на листе «Сводка».').font = f_sub

ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = f'{H1}:{HN}'
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)

# ================= Сводка =================
s2 = wb.create_sheet('Сводка')
for c, w in {'A': 26, 'B': 16, 'C': 12, 'D': 16, 'E': 13, 'F': 15, 'G': 18}.items():
    s2.column_dimensions[c].width = w
s2['A1'] = 'Сводка по кабелю'; s2['A1'].font = f_title
s2['A2'] = ('Все значения — формулы от листа «Кабельный журнал»: при изменении спусков, марки или сечения '
            'сводка пересчитается.')
s2['A2'].font = f_sub


def table_header(ws_, r, heads):
    for i, h in enumerate(heads):
        c = ws_.cell(r, i + 1, h)
        c.font = f_hdr; c.fill = fill_hdr; c.border = border; c.alignment = wrap_c
    ws_.row_dimensions[r].height = 32


def rng(col):
    return f"{JS}!${col}${FIRST_DATA}:${col}${LAST_DATA}"


s2['A4'] = '1. По марке и сечению кабеля'; s2['A4'].font = f_bold
table_header(s2, 5, ['Марка кабеля', 'Кол-во и сечение жил, мм²', 'Линий, шт.', 'Горизонтальная трасса, м',
                     'Спуски, м', 'Длина кабеля, м', 'С округлением вверх до 1 м, м'])
combos = []
for g in groups:
    k = (g['cable'], g['cores'])
    if k not in combos:
        combos.append(k)
r = 6
for cab, cor in combos:
    s2.cell(r, 1, cab); s2.cell(r, 2, cor)
    crit = f'{rng(C_BRAND)},A{r},{rng(C_CORES)},B{r}'
    s2.cell(r, 3, f'=COUNTIFS({crit})')
    s2.cell(r, 4, f'=SUMIFS({rng(C_H)},{crit})')
    s2.cell(r, 5, f'=SUMIFS({rng(C_DROPS)},{crit})')
    s2.cell(r, 6, f'=SUMIFS({rng(C_LEN)},{crit})')
    s2.cell(r, 7, f'=ROUNDUP(F{r},0)')
    r += 1
first_c, last_c = 6, r - 1
s2.cell(r, 1, 'Итого')
for col in 'CDEFG':
    s2[f'{col}{r}'] = f'=SUM({col}{first_c}:{col}{last_c})'
TOT2 = r
r += 1
s2.cell(r, 1, 'Контроль: ВСЕГО по журналу')
s2[f'F{r}'] = f"={JS}!{C_LEN}{TOT}"
CHK = r
r += 1
s2.cell(r, 1, 'Расхождение (должно быть 0)')
s2[f'F{r}'] = f'=ROUND(F{TOT2}-F{CHK},2)'
for rr in range(6, r + 1):
    for cc in range(1, 8):
        c = s2.cell(rr, cc)
        c.border = border
        c.font = f_sub if rr > TOT2 and cc == 1 else f_txt
        if cc >= 3:
            c.number_format = '0' if cc in (3, 7) else '0.00'
            c.alignment = right
    if rr == TOT2:
        for cc in range(1, 8):
            s2.cell(rr, cc).font = f_bold; s2.cell(rr, cc).fill = fill_tot
s2[f'F{CHK}'].font = f_link
for rr in range(6, TOT2):
    for cc in (1, 2):
        s2.cell(rr, cc).fill = fill_ass

r += 2
s2.cell(r, 1, '2. По системам').font = f_bold
r += 1
table_header(s2, r, ['Система', 'Линий, шт.', 'Горизонтальная трасса, м', 'Спуски, м', 'Длина кабеля, м'])
r += 1
sys_first = r
for sec in sections:
    sr = subtot_rows[sec]
    fr, lr = sec_bounds[sec]
    s2.cell(r, 1, sec)
    s2.cell(r, 2, f'=COUNT({JS}!{C_H}{fr}:{C_H}{lr})')
    s2.cell(r, 3, f'={JS}!{C_H}{sr}')
    s2.cell(r, 4, f'={JS}!{C_DROPS}{sr}')
    s2.cell(r, 5, f'={JS}!{C_LEN}{sr}')
    r += 1
s2.cell(r, 1, 'Итого')
for col in 'BCDE':
    s2[f'{col}{r}'] = f'=SUM({col}{sys_first}:{col}{r - 1})'
for rr in range(sys_first, r + 1):
    for cc in range(1, 6):
        c = s2.cell(rr, cc)
        c.border = border
        c.font = f_bold if rr == r else (f_link if cc > 1 else f_txt)
        if cc > 1:
            c.number_format = '0' if cc == 2 else '0.00'
            c.alignment = right
        if rr == r:
            c.fill = fill_tot

r += 2
s2.cell(r, 1, '3. Спуски').font = f_bold
r += 1
table_header(s2, r, ['Вид спуска', 'Кол-во, шт.', 'Всего, м'])
r += 1
sp_first = r
for lab, col, par in (('К щитку', C_NP, P_PANEL), ('К розеткам', C_NS, P_SOCK),
                      ('К выключателям', C_NW, P_SW), ('К светильникам', C_NL, P_LUM)):
    s2.cell(r, 1, lab)
    s2.cell(r, 2, f'={JS}!{col}{TOT}')
    # same rounding as in the journal rows, over data rows only (they have a number in column A)
    s2.cell(r, 3, f'=SUMPRODUCT(ISNUMBER({rng(C_NO)})*ROUND({rng(col)}*{JS}!{par}{FACTOR},2))')
    r += 1
s2.cell(r, 1, 'Итого')
s2[f'B{r}'] = f'=SUM(B{sp_first}:B{r - 1})'
s2[f'C{r}'] = f'=SUM(C{sp_first}:C{r - 1})'
for rr in range(sp_first, r + 1):
    for cc in range(1, 4):
        c = s2.cell(rr, cc)
        c.border = border
        c.font = f_bold if rr == r else (f_link if cc == 2 else f_txt)
        if cc > 1:
            c.number_format = '0' if cc == 2 else '0.00'
            c.alignment = right
        if rr == r:
            c.fill = fill_tot

# ================= Пояснения =================
s3 = wb.create_sheet('Пояснения')
s3.column_dimensions['A'].width = 4
s3.column_dimensions['B'].width = 120
s3['A1'] = 'Пояснения к кабельному журналу'; s3['A1'].font = f_title
r = 3
for block in notes:
    s3.cell(r, 1, block['title']).font = f_bold
    s3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1
    for i, line in enumerate(block['items'], 1):
        s3.cell(r, 1, f'{i}.').font = f_txt
        s3.cell(r, 1).alignment = Alignment(horizontal='right', vertical='top')
        c = s3.cell(r, 2, line); c.font = f_txt; c.alignment = Alignment(wrap_text=True, vertical='top')
        s3.row_dimensions[r].height = est_height([(line, 120)])
        r += 1
    r += 1

for sh in (s2, s3):
    sh.page_setup.orientation = 'landscape'
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print('saved', OUT, 'data rows', FIRST_DATA, '-', LAST_DATA, 'total row', TOT, 'subtotals', subtot_rows)
