"""Workbook with floor areas of rooms. Usage: python3 build_areas_xlsx.py out.xlsx"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from areas import room_polygons, ORDER
from config import DWG_JSON
from rooms import ROOM_NAMES

OUT = sys.argv[1]
db = json.load(open(DWG_JSON))
rooms, col_in, other = room_polygons(db)

F = 'Arial'
thin = Side(style='thin', color='808080')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='DCE6F1')
tot_fill = PatternFill('solid', fgColor='E2EFDA')

wb = Workbook()
ws = wb.active
ws.title = 'Площади помещений'
for c, w in {'A': 6, 'B': 26, 'C': 16, 'D': 62}.items():
    ws.column_dimensions[c].width = w
ws['A1'] = 'Площади полов помещений'
ws['A1'].font = Font(name=F, size=14, bold=True)
ws['A2'] = ('Посчитано по присланному чертежу DWG (план этажа): площадь внутри контура стен, без дверных проёмов. '
            'Названия помещений — по данным заказчика.')
ws['A2'].font = Font(name=F, size=9, italic=True, color='404040')
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:D2')
ws.row_dimensions[2].height = 26

H = 4
for i, h in enumerate(['№ п/п', 'Помещение', 'Площадь пола, м²', 'Примечание'], 1):
    c = ws.cell(H, i, h)
    c.font = Font(name=F, size=10, bold=True); c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[H].height = 30

r = H + 1
for i, n in enumerate(ORDER, 1):
    note = ''
    if col_in[n]:
        a = sum(c.area for c in col_in[n])
        note = f'за вычетом колонны 0,45×0,45 м ({a:.2f} м²)'.replace('.', ',')
    vals = [i, ROOM_NAMES[n], round(rooms[n].area, 2), note or None]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.border = border
        c.font = Font(name=F, size=10, color='0000FF' if j == 3 else '000000')
        c.alignment = Alignment(horizontal='center' if j == 1 else ('right' if j == 3 else 'left'), vertical='center')
        if j == 3:
            c.number_format = '0.00'
    r += 1
first, last = H + 1, r - 1
ws.cell(r, 2, 'Итого').font = Font(name=F, size=10, bold=True)
c = ws.cell(r, 3, f'=SUM(C{first}:C{last})')
c.font = Font(name=F, size=10, bold=True); c.number_format = '0.00'; c.alignment = Alignment(horizontal='right')
for j in range(1, 5):
    ws.cell(r, j).border = border; ws.cell(r, j).fill = tot_fill

notes = [
    'Площадь — по внутренним граням стен на плане. Проёмы дверей в площадь не входят.',
    'Колонна в кабинете 1022 исключена из площади пола.',
    'Пространства между санузлами на чертеже нарисованы сплошным массивом стены, без дверей (шахты) — не учтены.',
    'Синий шрифт — значения, снятые с чертежа; итог — формула.',
]
r += 2
ws.cell(r, 1, 'Пояснения').font = Font(name=F, size=10, bold=True)
for t in notes:
    r += 1
    c = ws.cell(r, 1, t); c.font = Font(name=F, size=9, color='404040')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print('saved', OUT, 'rows', first, '-', last)
